from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "工业数据分析与数据工程简历.docx"

BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(80, 80, 80)
ACCENT = RGBColor(27, 66, 98)


def set_run_font(run, name="Microsoft YaHei", size=9.5, bold=False, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1B4262")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def compact(paragraph, before=0, after=0, line=1.05):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    return paragraph


def section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    compact(p, before=5, after=2, line=1)
    r = p.add_run(text)
    set_run_font(r, size=10.5, bold=True, color=BLACK)


def add_bullet(doc, lead, body):
    p = doc.add_paragraph(style="List Bullet")
    compact(p, after=1.4, line=1.06)
    p.paragraph_format.left_indent = Cm(0.42)
    p.paragraph_format.first_line_indent = Cm(-0.22)
    set_run_font(p.add_run(lead), size=9.2, bold=True)
    set_run_font(p.add_run(body), size=9.2)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.05)
section.bottom_margin = Cm(0.9)
section.left_margin = Cm(1.45)
section.right_margin = Cm(1.45)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(9.5)
styles["Title"].font.color.rgb = BLACK
styles["Heading 1"].font.color.rgb = BLACK
title_ppr = styles["Title"]._element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)

# Header block
p = doc.add_paragraph(style="Title")
compact(p, after=0, line=1)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(p.add_run("姓名待填写"), size=19, bold=True, color=BLACK)
p_pr = p._p.get_or_add_pPr()
for border in p_pr.findall(qn("w:pBdr")):
    p_pr.remove(border)

p = doc.add_paragraph()
compact(p, after=2, line=1)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(p.add_run("数据分析师  数据工程师"), size=10.5, bold=True, color=ACCENT)

p = doc.add_paragraph()
compact(p, after=4, line=1)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(p.add_run("电话待填写  |  邮箱待填写  |  城市待填写  |  "), size=9, color=GRAY)
add_hyperlink(p, "GitHub", "https://github.com/YZC0219/industrial_energy_analysis")
set_run_font(p.add_run("  |  "), size=9, color=GRAY)
add_hyperlink(p, "在线演示", "https://yzc0219.github.io/industrial_energy_analysis/")

section_heading(doc, "个人概况")
p = doc.add_paragraph()
compact(p, after=1, line=1.08)
set_run_font(p.add_run("具备从业务问题定义、数据质量治理、维度建模、SQL 分析到调度与可视化交付的端到端项目经验。"), size=9.3)
set_run_font(p.add_run("擅长把工业场景的物理边界和统计口径落实为可追溯的数据模型、异常检测方法与自动化测试。"), size=9.3)

section_heading(doc, "核心技能")
table = doc.add_table(rows=2, cols=2)
table.autofit = False
table.columns[0].width = Cm(9.0)
table.columns[1].width = Cm(9.0)
remove_table_borders(table)
skills = [
    ("数据分析", "SQL  窗口函数  同比环比  单耗  帕累托  碳排核算"),
    ("数据工程", "Python  pandas  NumPy  MySQL  Hive  Spark SQL  DataX"),
    ("建模与检测", "星型模型  2σ  最小二乘基线  CUSUM  蒙特卡洛校准"),
    ("工程化", "Airflow  Docker Compose  pytest  GitHub Actions  HTML CSS JS SVG"),
]
for i, (label, value) in enumerate(skills):
    cell = table.cell(i // 2, i % 2)
    set_cell_margins(cell, 25, 25, 25, 25)
    p = cell.paragraphs[0]
    compact(p, line=1)
    set_run_font(p.add_run(f"{label}："), size=8.8, bold=True, color=ACCENT)
    set_run_font(p.add_run(value), size=8.8)

section_heading(doc, "项目经历")
p = doc.add_paragraph()
compact(p, after=0, line=1)
p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
set_run_font(p.add_run("工业能耗数据仓库与异常检测系统"), size=10.3, bold=True)
set_run_font(p.add_run("\t2026.09"), size=9, color=GRAY)
p = doc.add_paragraph()
compact(p, after=2, line=1)
set_run_font(p.add_run("个人项目  |  本科毕业设计  |  可复现模拟数据：8 个车间、6 种能源、731 天"), size=8.8, color=GRAY)

add_bullet(doc, "端到端交付：", "独立完成模拟脏数据生成、可追溯清洗、MySQL 星型模型、29 组 SQL 分析、异常检测、Airflow 调度及自包含交互式报告，并通过 GitHub Actions 自动发布。")
add_bullet(doc, "数据质量闭环：", "主动注入并治理 10 类质量问题；从 CUSUM 异常信号追溯到“离群值整行删除”缺陷，将能源品种缺失的车间日由 34 个降至 0，并新增回归测试固化修复。")
add_bullet(doc, "业务口径设计：", "建立全厂、剔除公用工程、仅连续型车间三类口径，识别动力站二次能源重复计量占全厂能耗 27.23%，避免单耗和相关性分析失真。")
add_bullet(doc, "异常检测：", "实现 2σ、考虑产量与季节项的基线模型及双侧 CUSUM；通过 5,000 次白噪声仿真，将长序列判定限从经验值 5σ 重新标定为 9.9σ。")
add_bullet(doc, "业务价值量化：", "识别 4 个间歇生产车间 58 个停产日，估算待机损耗约 79.39 万元；按 2024 年单耗基线测算 2025 年节能约 731 tce、减排约 1,797 tCO₂、能源毛收益约 247 万元。")
add_bullet(doc, "工程可靠性：", "设计幂等 upsert、水位线增量装载、失败重跑、查询快照与端到端验证；当前本地自动化检查 82 项通过，并以恒等式校验前端筛选聚合结果。")
add_bullet(doc, "架构扩展：", "完成 Hive ODS DWD DWS ADS 分层、DataX 全量与增量同步、Spark SQL 分层处理及 Airflow 两级质量门禁的代码与离线契约验收；明确标注真实集群性能测试尚待部署环境执行。")
add_bullet(doc, "预测基线：", "构建无泄漏特征管道与滚动时序验证框架，完成 2,928 条 7 日季节性朴素预测，MAE 0.7513 tce、RMSE 1.3144 tce，并在 CI 中自动复算与留存证据。")

section_heading(doc, "教育经历")
p = doc.add_paragraph()
compact(p, after=0, line=1)
p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
set_run_font(p.add_run("学校待填写  |  专业待填写  |  本科"), size=9.5, bold=True)
set_run_font(p.add_run("\t毕业时间待填写"), size=9, color=GRAY)
p = doc.add_paragraph()
compact(p, after=0, line=1)
set_run_font(p.add_run("毕业论文：工业能耗数据仓库与异常检测系统的设计与实现"), size=9.1)

# Document metadata
doc.core_properties.title = "数据分析与数据工程求职简历"
doc.core_properties.subject = "基于工业能耗数据仓库与异常检测项目成果"
doc.core_properties.author = "姓名待填写"
doc.core_properties.keywords = "数据分析, 数据工程, Python, SQL, MySQL, Airflow, 工业能耗"

doc.save(OUT)
print(OUT)
