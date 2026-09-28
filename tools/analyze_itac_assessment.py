"""Build a reproducible real-assessment case from the public ITAC workbook.

Requires openpyxl (see requirements-itac.txt). The source workbook stays
separate from the project's simulated MES/ERP energy data.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
from datetime import datetime, timezone
from pathlib import Path

os.environ["OPENBLAS_NUM_THREADS"] = "1"
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "data" / "real" / "itac_2026" / "ITAC_Database_20260927.xlsx"
DEFAULT_OUTPUT = ROOT / "output" / "phase5_itac"
CASE_ID = "DL0238"
DESCRIPTIONS = {
    "2.4226": "Use / purchase optimum sized compressor",
    "2.4224": "Upgrade controls on compressors",
    "2.4239": "Eliminate or reduce compressed air usage",
    "2.2113": "Repair or replace steam traps",
    "2.2437": "Recover waste heat from equipment",
    "2.2135": "Repair and eliminate steam leaks",
    "2.1392": "Replace purchased steam with other energy source",
}


def sha256(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


def rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    header = next(iterator)
    for values in iterator:
        if any(value is not None for value in values):
            yield dict(zip(header, values))


def number(value) -> float:
    if value is None or value == "":
        return 0.0
    return float(value)


def code(value) -> str:
    return f"{float(value):.4f}" if value is not None else ""


def analyze(input_path: Path, output_dir: Path, assessment_id: str = CASE_ID) -> dict:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    assessment = next((row for row in rows(workbook["ASSESS"])
                       if row["ID"] == assessment_id), None)
    if assessment is None:
        raise ValueError(f"找不到评估编号 {assessment_id}")
    recommendations = [row for row in rows(workbook["RECC"])
                       if row["ID"] == assessment_id]
    if len(recommendations) != int(assessment["RECC_COUNT_recc"]):
        raise ValueError("评估表的建议数量与建议表不一致")
    item_rows = []
    for row in recommendations:
        arc = code(row["ARC2"])
        annual_savings = sum(number(row[field]) for field in
                             ("PSAVED", "SSAVED", "TSAVED", "QSAVED"))
        cost = number(row["IMPCOST"])
        item_rows.append({
            "recommendation_id": row["SUPERID"],
            "arc_code": arc,
            "description": DESCRIPTIONS.get(arc, "Consult ITAC ARC code index"),
            "implementation_status": row["IMPSTATUS"] or "unknown",
            "estimated_annual_savings_usd": annual_savings,
            "estimated_implementation_cost_usd": cost,
            "estimated_simple_payback_years": cost / annual_savings if annual_savings > 0 else None,
            "primary_resource_code": row["PSOURCCODE"],
            "primary_resource_conserved": row["PCONSERVED"],
            "primary_cost_savings_usd": row["PSAVED"],
            "secondary_resource_code": row["SSOURCCODE"],
            "secondary_resource_conserved": row["SCONSERVED"],
            "secondary_cost_savings_usd": row["SSAVED"],
            "source_page": f"https://itac.university/assessment/{assessment_id}",
        })
    item_rows.sort(key=lambda row: row["recommendation_id"])
    known = [row for row in item_rows if row["implementation_status"] in ("I", "N")]
    implemented = [row for row in known if row["implementation_status"] == "I"]
    electricity_kwh = number(assessment["EC_plant_usage"])
    annual_production = number(assessment["PRODLEVEL"])
    result = {
        "assessment_id": assessment_id,
        "fiscal_year": assessment["FY"],
        "state": assessment["STATE"],
        "naics": assessment["NAICS"],
        "product": assessment["PRODUCTS"],
        "annual_production": annual_production,
        "production_unit_code": assessment["PRODUNITS"],
        "production_unit_for_case": "ton (confirmed on official assessment page)" if assessment_id == CASE_ID else None,
        "annual_electricity_kwh": electricity_kwh,
        "annual_electricity_usage_cost_usd": number(assessment["EC_plant_cost"]),
        "electricity_kwh_per_production_unit": electricity_kwh / annual_production if annual_production > 0 else None,
        "recommendations": len(item_rows),
        "known_implementation_status": len(known),
        "implemented_status_count": len(implemented),
        "estimated_recommended_annual_savings_usd": sum(row["estimated_annual_savings_usd"] for row in item_rows),
        "estimated_implemented_annual_savings_usd": sum(row["estimated_annual_savings_usd"] for row in implemented),
        "estimated_implemented_cost_usd": sum(row["estimated_implementation_cost_usd"] for row in implemented),
        "source_page": f"https://itac.university/assessment/{assessment_id}",
        "source_workbook_sha256": sha256(input_path),
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "limit": "Savings and costs are ITAC engineering estimates. I means the measure was reported implemented; this workbook has no metered post-intervention consumption or dated before/after series.",
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "assessment_summary.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    with (output_dir / "recommendations.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=item_rows[0].keys())
        writer.writeheader()
        writer.writerows(item_rows)
    _write_report(output_dir / "assessment_case.md", result, item_rows)
    return result


def _write_report(path: Path, summary: dict, items: list[dict]) -> None:
    production_unit = "ton（源数据单位，未换算）" if summary["assessment_id"] == CASE_ID else "原始产量单位代码"
    intensity_unit = "ton" if summary["assessment_id"] == CASE_ID else "原始产量单位代码"
    lines = [
        f"# ITAC 真实工厂评估案例：{summary['assessment_id']}",
        "",
        f"来源：[美国能源部 ITAC 官方评估页]({summary['source_page']})，{summary['fiscal_year']} 财年。",
        "",
        "## 工厂基线",
        "",
        f"- 产品：{summary['product']}；年产量：{summary['annual_production']:,.0f} {production_unit}。",
        f"- 年用电量：{summary['annual_electricity_kwh']:,.0f} kWh；用电成本：${summary['annual_electricity_usage_cost_usd']:,.0f}。",
        f"- 年电耗强度：{summary['electricity_kwh_per_production_unit']:,.2f} kWh/{intensity_unit}。这是单个年度的基线，不能由此推出同比节能。",
        "",
        "## 建议与记录的实施状态",
        "",
        "| 建议 | 措施 | 状态 | 年节省估计 (USD) | 投资估计 (USD) | 简单回收期 (年) |",
        "|---|---|---|---:|---:|---:|",
    ]
    for item in items:
        payback = item["estimated_simple_payback_years"]
        lines.append(f"| {item['recommendation_id']} | {item['description']} | {item['implementation_status']} | "
                     f"{item['estimated_annual_savings_usd']:,.0f} | {item['estimated_implementation_cost_usd']:,.0f} | "
                     f"{payback:.2f} |" if payback is not None else
                     f"| {item['recommendation_id']} | {item['description']} | {item['implementation_status']} | "
                     f"{item['estimated_annual_savings_usd']:,.0f} | {item['estimated_implementation_cost_usd']:,.0f} | — |")
    lines.extend([
        "", f"数据库记录 {summary['implemented_status_count']}/{summary['known_implementation_status']} 项建议已实施。",
        f"这些已实施建议的年节省估计合计 ${summary['estimated_implemented_annual_savings_usd']:,.0f}，投资估计合计 ${summary['estimated_implemented_cost_usd']:,.0f}。",
        "",
        "## 解释边界",
        "",
        "- `I` 是 ITAC 数据库记录的实施状态；年节省和成本是评估方工程估计。工作簿没有实施前后连续计量序列，不能把估计额写成实测节能或因果效果。",
        "- 该工厂的电耗强度仅与本厂同一产品和同一产量单位可比。ITAC 与 UCI 钢厂数据属于不同企业，不能按时间或设备键拼接。",
        "- 回收期为投资估计除以年节省估计，未计融资、维护、价格变动、产能影响和税费。",
        "", "来源工作簿哈希和逐项资源明细分别见 `assessment_summary.json`、`recommendations.csv`。", "",
    ])
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--assessment-id", default=CASE_ID)
    args = parser.parse_args()
    print(json.dumps(analyze(args.input, args.output_dir, args.assessment_id),
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
