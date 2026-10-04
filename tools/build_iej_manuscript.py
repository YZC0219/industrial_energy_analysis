from pathlib import Path
import json, re
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output'/'paper'
REAL=OUT/'uci_steel'
DEST=ROOT/'IEJ_industrial_energy_forecasting_manuscript.docx'
syn=json.loads((OUT/'paper_experiment_summary.json').read_text(encoding='utf-8'))
uci=json.loads((REAL/'uci_steel_experiment_summary.json').read_text(encoding='utf-8'))
metrics=json.loads((ROOT/'output'/'ml_model_metrics.json').read_text(encoding='utf-8'))
deep=json.loads((ROOT/'output'/'ml_deep_metrics.json').read_text(encoding='utf-8'))

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.27); sec.bottom_margin=Cm(1.27); sec.left_margin=Cm(2.54); sec.right_margin=Cm(1.27)
sectPr=sec._sectPr
cols=OxmlElement('w:cols'); cols.set(qn('w:num'),'2'); cols.set(qn('w:space'),'333'); sectPr.append(cols)
styles=doc.styles
for s in ['Normal','Body Text']:
 st=styles[s]; st.font.name='Times New Roman'; st.font.size=Pt(10)
 st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
 st.paragraph_format.line_spacing=1.05; st.paragraph_format.space_after=Pt(3)
for name,size in [('Title',15),('Heading 1',11),('Heading 2',10)]:
 st=styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=None
 st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
 st.paragraph_format.keep_with_next=True; st.paragraph_format.space_before=Pt(5); st.paragraph_format.space_after=Pt(2)

def para(t='',style=None,bold=False,align=None,italic=False):
 p=doc.add_paragraph(style=style)
 if align is not None:p.alignment=align
 r=p.add_run(t);r.bold=bold;r.italic=italic
 return p
def head(t,l=1): doc.add_heading(t,level=l)
def shade(cell,color='D9E2F3'):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd');shd.set(qn('w:fill'),color);tcPr.append(shd)
def table(caption,headers,rows):
 para(caption,bold=True)
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 col_width=Cm(7.8/len(headers))
 tbl_w=t._tbl.tblPr.find(qn('w:tblW'));tbl_w.set(qn('w:w'),str(round(Cm(7.8).twips)));tbl_w.set(qn('w:type'),'dxa')
 for grid_col in t._tbl.tblGrid.gridCol_lst:grid_col.set(qn('w:w'),str(round(col_width.twips)))
 for column in t.columns:column.width=col_width
 for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
 for cell in t.rows[0].cells:cell.width=col_width
 trPr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');rep.set(qn('w:val'),'true');trPr.append(rep)
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v);cells[i].width=col_width
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(8)
 return t
def fig(path,caption,width=8.0):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 p.add_run().add_picture(str(path),width=Cm(width))
 para(caption,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)

# Front matter
p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Rolling Validation and Explainable Forecasting of Industrial Energy Use')
para('Yang Zicheng¹',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
para('¹Chuzhou University, No. 1 West Huifeng Road, Chuzhou, Anhui 239000, China',align=WD_ALIGN_PARAGRAPH.CENTER)
para('Corresponding author: Yang Zicheng; E-mail: yangzicheng02190531@gmail.com',align=WD_ALIGN_PARAGRAPH.CENTER)
head('Abstract',1)
abstract=('Industrial energy forecasts are difficult to compare when time ordering, feature availability and evaluation windows are not aligned. This study evaluates a reproducible next-day forecasting workflow using two distinct data settings: a fixed-seed synthetic panel representing eight workshops and a public steel-facility electricity series. On the synthetic panel, 11 expanding-window folds yielded 2,640 workshop-day predictions. LightGBM achieved a pooled mean absolute error (MAE) of 0.4674 project-defined coal-equivalent resource-index units, compared with 0.7550 for a seven-day seasonal-naive baseline, 0.5213 for a long short-term memory network and 0.5448 for a Transformer. Date-paired ablations showed higher mean MAE after removing weather/calendar features (0.0569 units; moving-block 95% interval 0.0389–0.0769 for 14-day blocks) or target history (0.0508; 0.0180–0.0793); these are exploratory estimates from one generated panel. TreeSHAP assigned the largest mean absolute contributions to recent energy lags. In a separate seven-fold evaluation of 365 daily totals aggregated from 35,040 quarter-hour observations at one steel facility, LightGBM achieved an MAE of 769.9 kWh/day, compared with 869.6 kWh/day for the seasonal-naive baseline. Removing recent history increased daily MAE by 152.0 kWh/day (14-day block interval 78.7–216.6). The synthetic target uses project conversion factors and is not a standards-compliant measure of total factory energy. Because the synthetic outcomes reflect programmed rules and the public evaluation covers one site and one year, the results establish a reproducible analysis workflow rather than general industrial performance or causal effects.')
para(abstract)
para('Keywords: Energy forecasting; Feature ablation; Industrial energy; SHAP; Time-series validation.',bold=True)

head('1. Introduction',1)
paras=[
"Energy forecasting supports production planning, procurement, load management and operational monitoring. In manufacturing settings, however, the term energy consumption can refer to different targets: machine-level electricity, process energy, total factory demand or energy normalized by output. A model evaluated at one boundary may not answer a question posed at another. Reviews of manufacturing energy prediction document substantial variation in system boundaries, input variables, temporal resolution and forecast horizons, which complicates comparison across published results [1,2]. Industrial energy and carbon forecasting also face uneven measurement quality and limited access to process variables, placing constraints on transfer between facilities [3].",
"Prior work spans process-level and plant-level models, smart-factory data mining, production planning and deep learning [4-10]. Studies have examined manufacturing energy modelling with machine learning [4], small and medium enterprise forecasting [5], process energy model selection [6], production-state profiling [7], predictive planning in energy-intensive industries [8], and smart-factory prediction [9]. Research in mining and gas operations has also considered interpretable forecasts [10,11]. These studies demonstrate that useful predictors and suitable model classes depend on the task and data boundary. Yet many comparisons remain difficult to interpret when training and test windows differ, future information is not explicitly audited, or the test period represents only one operating regime.",
"The literature also reflects different definitions of prediction success. A plant-level energy forecast may support procurement or capacity planning, while a process-level forecast may support scheduling and equipment coordination. Energy profiling methods can include production-state prediction and production estimation as intermediate tasks [7]. Work on predictive planning links energy information with production decisions [8,15], whereas a process model selection study may prioritize fidelity to a particular manufacturing sequence [6]. These outcomes are related but not interchangeable. The present study selects a single, measurable target: next-day aggregate energy. It does not infer energy intensity, operational efficiency, process state or optimal schedules from that target.",
"Short-term forecasting studies in process industries have used hybrid neural methods and operational covariates [12]. Broader reviews of manufacturing processes describe model choices ranging from engineering calculations to statistical learning [13,14]. Predictive scheduling and resource allocation add another connection between energy forecasts and production decisions [15]. Research on industrial building energy has evaluated statistical and machine-learning approaches, but its occupancy, building-envelope and HVAC context differs from process manufacturing [16]. Multi-scale manufacturing methods likewise emphasize that predictions depend on the aggregation level [17]. These strands motivate a study that defines the target and prediction-time information before comparing models.",
"Tree ensembles are common for structured tabular inputs, while recurrent networks and attention architectures are designed to represent sequential dependence. LightGBM provides an efficient gradient-boosted decision-tree implementation [18]; XGBoost offers a related scalable boosting framework [19]. Long short-term memory (LSTM) networks were developed to address long-range dependencies in recurrent learning [20], and the Transformer architecture uses attention rather than recurrence [21]. Model flexibility alone does not establish better out-of-sample forecasts. A strong seasonal-naive baseline and temporally ordered evaluation remain necessary, particularly where the series is short or seasonal.",
"This paper addresses a practical methodological gap: a forecast comparison should preserve time order, use only information available at the forecast origin, assess feature groups without changing the test cases, quantify uncertainty at the temporal replication level, and distinguish model explanation from causal evidence. Rolling-origin evaluation is a standard way to expose changes across forecast periods [22,23]. Predictive accuracy comparisons have formal statistical treatments [24], but a single dataset often supplies few independent time blocks. For this reason, the current study pairs daily losses and uses moving-block intervals; folds define model-refitting periods rather than independent observations.",
"The study asks three questions. First, how do a seasonal-naive baseline, LightGBM, LSTM and Transformer compare under a common rolling protocol on a controlled synthetic workshop panel? Second, which prespecified feature groups change LightGBM fold-level MAE in that panel, and are the changes consistent across time folds? Third, do model attributions from TreeSHAP align with the feature-group findings, and what changes when the workflow is applied to a public real-measurement series from a steel facility? SHAP provides additive feature attributions for model predictions, but its values describe the fitted model and the chosen data distribution rather than interventions or physical causality [25]. Reviews from building-energy forecasting provide complementary evidence about data-driven methods but address different physical systems, so their findings require care when transferred to manufacturing [26-28].",
"The contribution is a documented, reproducible workflow with a controlled synthetic experiment and a separately reported single-facility replication. The synthetic data provide a setting in which missingness and operational factors can be generated consistently, making the protocol auditable; they do not supply evidence of real plant performance. The public-data application is used to test whether the workflow can be applied to observed data, not to claim external validity across manufacturing. Throughout the paper, synthetic and measured results are kept separate, and operational recommendations are limited to what these data support."
]
for x in paras:para(x)

head('2. Data and study design',1)
head('2.1 Synthetic workshop panel',2)
for x in [
"The primary controlled experiment uses a fixed-seed synthetic panel generated by the project workflow (seed 20240918). It contains eight workshops, six energy/resource carriers and 731 daily dates from 1 January 2024 through 31 December 2025. The cleaned target is a project-defined coal-equivalent resource index in tonnes-equivalent units, calculated by summing carrier quantities multiplied by the project factors in Table 2. It is not a standards-compliant measure of total factory energy: water and compressed air are included as resource carriers, and the water factor differs from the equivalent value listed in GB/T 2589-2020. The factors are retained to reproduce the existing synthetic experiment, not endorsed as a universal conversion boundary. The data generator encodes annual production growth, declining unit energy intensity, seasonal temperature, production schedules, holidays, energy-price variation and operating-load states. Random disturbances are added to weather, energy use, prices and production. Missing fields, aliases and malformed values are also injected before cleaning. These operations test the end-to-end pipeline, but the underlying signal is intentionally constructed and cannot stand in for an observed industrial population.",
"The panel permits multiple workshop series to share calendar conditions while retaining distinct workshop identities. This resembles a repeated measurement structure, but it should not be confused with a random sample of eight actual factories: workshops are generated entities and their dependencies follow the simulator. The target is an aggregate daily resource index, not energy intensity per unit of output. A lower forecast error therefore does not imply improved efficiency. No intervention, energy-saving action or fault label is present in the experiment."
,
"The synthetic records pass through a cleaning pipeline before modelling. The cleaned feature table contains 5,848 workshop-day rows (eight workshops by 731 dates); 5,624 rows meet feature eligibility after lag warm-up, and the 11 test folds contain 2,640 predictions per model. The simulator creates energy-detail rows, after which identifiers, aliases, dates and numeric measures are normalized and aggregated to the workshop-day target. Deliberately malformed values and missing entries challenge the quality checks, while the analysis uses the resulting cleaned daily series. These generated quality problems are reproducible stressors rather than a claim that the data have the same error distribution as plant meters. The model sees cleaned features, and the paper evaluates prediction after that preparation step rather than evaluating a joint imputation and forecasting system.",
"The simulated annual trend and unit-intensity change produce gradual shifts in consumption, while schedules, holidays and temperature create recurring patterns. Perturbations add noise around these components. The panel gives the model several information paths: persistence from energy lags, predictable calendar context, and lagged operational covariates. The purpose is to observe whether the workflow recovers known structural regularities. It is not a blinded synthetic benchmark because the researcher knows how the generator was constructed, and findings are interpreted with that knowledge."
]:para(x)
head('2.2 Public steel-facility series',2)
for x in [
"The external measured-data analysis uses the Steel Industry Energy Consumption dataset in the UCI Machine Learning Repository [29]. The dataset records electricity use at 15-minute intervals for a steel facility operated by DAEWOO Steel Co. in Gwangyang, Republic of Korea, during 2018. It contains 35,040 records and is distributed under the Creative Commons Attribution 4.0 license. The target, Usage_kWh, was summed over 96 quarter-hour observations per calendar day to create 365 daily totals. After timestamp ordering, the interval sequence was complete, with no duplicate timestamps, missing 15-minute intervals or missing target values.",
"The dataset is not a multi-factory panel. It does not provide measured production quantities, workshop identifiers or sufficient process metadata to explain all operating changes. Several fields record contemporaneous power factors, reactive power, carbon emissions or load categories. Such same-day records could be unavailable when a next-day forecast is issued and can be algebraically or operationally connected to the target. To avoid using target-day post-outcome information, the forecasting inputs were restricted to past daily energy and known calendar variables. The real-data experiment is consequently narrower than the synthetic study and is not a direct replication of its feature set or training window."
]:para(x)
table('Table 1. Data settings and evaluation scope',['Setting','Data and evaluation'],[
['Synthetic panel','8 workshops × 731 days; 5,848 rows, 5,624 feature-eligible; target: project resource index; 11 folds, 2,640 predictions'],
['UCI steel facility','35,040 quarter-hour records aggregated to 365 daily kWh totals; 7 folds, 210 test days']])
table('Table 2. Project conversion factors used to construct the synthetic target',['Carrier and input unit','Project factor'],[
['Electricity (kWh)','0.1229 kgce/kWh'],['Natural gas (m³)','1.3300 kgce/m³'],['Steam (t)','95.7000 kgce/t'],['Industrial water (m³)','0.2429 kgce/m³'],['Compressed air (m³)','0.0400 kgce/m³'],['Diesel (kg)','1.4571 kgce/kg']])
para('The coefficients are the exact carrier factors implemented in the feature pipeline. GB/T 2589-2020 provides the reference context for standard coal-equivalent accounting, but the combined index above should not be interpreted as a standard-defined total-energy statistic. A standards-aligned field study must define its energy boundary and verify each carrier factor against the applicable measurement and accounting basis.',italic=True)

head('3. Forecasting methods',1)
head('3.1 Forecast origin and feature availability',2)
for x in [
"The forecast task is a rolling one-step-ahead prediction. For each series and date t, the model predicts energy use on day t using only information assumed to be available before that day begins. Calendar encodings for the target date can be known in advance. Measurements whose values are observed only during or after day t are shifted to t−1. For the synthetic panel, lagged predictors include production quantity, average temperature, low-load share, shutdown share and energy-price index. The energy target history includes lags and rolling summaries, with every rolling window shifted so that the target date is excluded.",
"This availability rule is more important than the variable name. A scheduled production quantity may be known at the origin, whereas realized output may not be. A weather observation differs from a weather forecast; an observed price may differ from a contracted tariff. In the present synthetic analysis, lagging realized measurements gives a conservative and reproducible rule. In operational reuse, each field would need an explicit timestamp and provenance label. The paper does not assume that a model can access same-day measured values merely because they exist in a retrospective dataset.",
"The rolling protocol also separates fitting from forecasting. At the beginning of each 30-day test block, the learner is fitted on data available up to that origin. During the block, each target is evaluated as a single-day forecast and lagged target features are built from observed history available before that target. This is a rolling one-step setup, rather than a fixed-origin 30-day trajectory in which all forecasts must be issued simultaneously. It corresponds to a daily operational cycle in which yesterday’s observation becomes available for today’s forecast. A use case that issues an entire month of forecasts at once would need a different recursive or direct multi-horizon design.",
"The synthetic feature set includes t−1, t−7 and t−14 energy lags; seven-day and 28-day rolling energy means; lagged production, temperature, operating-state shares and price index; day of week, weekend and annual sine/cosine encodings; and workshop identity indicators. The public steel series uses daily energy lags at 1, 4, 7 and 28 days, seven-day and 28-day means, day-of-week and annual calendar encodings, and a weekend flag. This second design reflects the smaller record and limited metadata. No imputation or feature scaling is allowed to use future test values."
]:para(x)
head('3.2 Benchmark models',2)
for x in [
"The synthetic comparison includes a seven-day seasonal-naive forecast, LightGBM, an LSTM and a Transformer. The seasonal-naive estimate uses the previous observation at the same weekday position, y(t−7), and provides a transparent reference for weekly persistence. LightGBM is fitted to the tabular lag and covariate representation, using 180 trees, learning rate 0.04, 31 leaves and maximum depth 7. The pseudorandom seed and project defaults are held fixed. The objective is regression; hyperparameters are not retuned separately on each test period.",
"The LSTM receives a 28-day sequence and a 24-unit hidden representation, trained for 12 epochs under the recorded project settings. The Transformer uses a 24-dimensional representation, four attention heads, one encoder layer and 12 training epochs. Both use the same target dates and rolling folds as the tree model. These architectures are comparative implementations rather than exhaustive architecture searches. The benchmark does not claim equal computational budgets, optimal tuning or statistical superiority of one model family. Its purpose is to establish performance under a common data and forecast protocol.",
"For the UCI series, the comparison uses LightGBM with 100 trees, seven leaves, maximum depth 3 and minimum child sample count 5, against the seven-day seasonal-naive baseline. The smaller initial training window and the restricted real-data features reflect the one-year record. Reusing the synthetic model configuration mechanically would ignore differences in sample size and feature availability; the distinction is recorded to prevent a false impression that both experiments use an identical protocol."
]:para(x)
para('Implementation used Python 3.13 with pandas 3.0.6, NumPy 2.5.3, LightGBM 4.7.0, SHAP 0.52.0 and PyTorch 2.14.0+cpu. The main analysis entry points are ml/paper_experiments.py and ml/uci_steel_experiment.py; the latter reads the cited UCI CSV. The synthetic feature pipeline is ml/feature_pipeline.py. The repository contains the scripts, while generated experiment summaries and intermediate predictions are build outputs.',italic=True)
head('3.3 Rolling-origin protocol and outcome metrics',2)
for x in [
"For the synthetic panel, 28 dates are reserved for feature warm-up. The initial training period contains 365 dates, and each test fold contains the next 30 dates. The training window expands as the origin advances by 30 days. Eleven complete folds are available, with eight workshops over 30 test days per fold, giving 240 records per fold and 2,640 predictions per model. No random shuffle is performed. The shared test rows ensure that model comparisons do not benefit from differing eligibility rules.",
"For the public steel series, 28 warm-up days are followed by an initial 120-day training period, a 30-day test block and 30-day increments. This produces seven folds and 210 daily test values per model. The initial training window is shorter than the synthetic window because only one year of observations is available. The two settings therefore answer related, not identical, questions.",
"The primary metrics are mean absolute error (MAE) and root mean squared error (RMSE). For n test values, MAE is the mean absolute difference between observed and predicted values, and RMSE is the square root of the mean squared difference. The pooled metrics combine all test predictions within a setting. Fold-level MAEs are reported descriptively. The synthetic pooled RMSE is not interchangeable with the unweighted mean of fold RMSE values; both are explicitly labelled when presented. Units are project-defined tonnes-equivalent for the synthetic resource index and kWh/day for the public steel series."
]:para(x)
head('3.4 Feature ablation and statistical inference',2)
for x in [
"Four synthetic feature groups were prespecified: production; weather and calendar; price and operating status; and target history. Each ablation removes one group while retaining all other inputs, training dates, test rows, and model settings. The UCI ablations remove calendar variables, longer monthly history, or recent energy history. The groups are designed to ask whether information sets improve the fitted forecast under the chosen protocol. They do not identify causal effects of production, weather, price, status or energy use.",
"For each test date, we first average absolute errors across the eight synthetic workshops, then define the paired difference as ablated-model daily MAE minus full-model daily MAE. The UCI series has one daily value. We report the mean daily difference and percentile intervals from 20,000 non-circular moving-block bootstrap resamples using 7-, 14- and 30-day blocks. This preserves short-range dependence within sampled blocks and makes sensitivity to block length visible; days are not treated as independent replicates.",
"These intervals remain exploratory. The 330 synthetic paired dates come from one generated panel, and the 210 UCI test dates come from one facility-year. Moving-block resampling cannot create independent sites or operating years, and its interval depends on block length and the observed test period. We therefore make no confirmatory significance claims and use effect estimates and interval sensitivity descriptively. Longer independent facility records are needed for stronger inference."
,
"Pairing compares the full and ablated models on the same dates, while block resampling retains local temporal dependence. Neither choice removes shared-period effects or establishes that the observed sequence represents other facilities or future years. A multi-facility study could use a hierarchical design that separates plant, period and day variation.",
"MAE expresses typical absolute error in the target unit and is less dominated by a small number of large residuals than RMSE. RMSE remains useful because large misses may matter for planning. Neither metric should be interpreted without the typical series scale, the cost of errors and a baseline."
]:para(x)
head('3.5 SHAP analysis',2)
for x in [
"TreeSHAP values were calculated for the LightGBM models. In the synthetic experiment, the first test-fold model was explained on a fixed-seed sample of 60 workshop-day cases. In the UCI analysis, all 30 cases in the first test fold were used. Global importance is summarized as the mean absolute SHAP value for each feature. The signed mean contribution is also retained in the underlying output but is not used to infer the direction of a physical effect.",
"A SHAP value describes how a feature contributes to a model prediction relative to the explainer’s expected output under its chosen background distribution. Correlated lag features may share or redistribute attribution, and importance can vary with period, training data and model specification [25]. We therefore read the SHAP ranking together with grouped ablation, rather than as an independent test of feature necessity. Attribution does not show what would happen if an operator intervened on a variable."
]:para(x)

head('4. Results',1)
head('4.1 Synthetic model comparison',2)
modelrows=[]
for m in metrics['models']+deep['models']:
 modelrows.append([m['model'],f"{m['mae']:.4f}",f"{m['rmse']:.4f}",str(m['samples'])])
table('Table 3. Synthetic-panel pooled forecast performance',['Model','Pooled MAE / RMSE; test cases'],[[r[0],f"{r[1]} / {r[2]}; n={r[3]}"] for r in modelrows])
for x in [
"LightGBM produced the lowest pooled MAE, 0.4674 project-defined tonnes-equivalent, against 0.7550 for the seasonal-naive model. The difference corresponds to a 38.1% lower MAE relative to that baseline in this generated panel. LSTM and Transformer MAEs were 0.5213 and 0.5448. The pooled RMSE values were 0.7750 for LightGBM, 1.3300 for the seasonal-naive forecast, 0.8538 for LSTM and 0.8415 for the Transformer. Thus, the Transformer had a slightly lower pooled RMSE than LSTM despite a higher MAE. This difference indicates sensitivity to larger residuals and should not be collapsed into a single ranking.",
"Performance varied across the 11 folds. The temporal variation matters because the synthetic series contains programmed seasonality, operating schedules and a gradual efficiency trend. A model that is best over the pooled test set may not be best in every period. In practical use, fold-level or month-level results help identify periods of degradation that an overall average can conceal. Here, however, the variability is generated by the simulator and cannot be interpreted as plant-level robustness."
,
"The comparison also shows why MAE and RMSE should be reported together. On the synthetic data the Transformer’s RMSE was slightly lower than the LSTM’s, while its MAE was higher. This pattern can occur when one model has fewer or smaller extreme residuals but somewhat larger typical errors, although the metric pair alone does not diagnose which cases caused the difference. A residual-by-time plot and subgroup breakdown could identify those cases. They were not used here to select a model, so the paper does not give a post hoc explanation for the relative RMSE ordering. The fold plot instead makes clear that temporal stage can affect the ranking.",
"The 38.1% relative MAE decrease compared with the naive reference is substantial within the generated panel, but a percentage improvement depends on the strength of that reference. If the baseline is weak because the series changes regime, relative improvement can look large; if the baseline already captures stable weekly structure, a smaller gain may still matter. A facility study should report operational thresholds, such as whether error is acceptable for a procurement or scheduling decision. Those thresholds are facility-specific and were not available in either dataset."
]:para(x)
fig(OUT/'figures'/'fold_mae.png','Figure 1. Fold-level MAE for synthetic-panel benchmarks.',width=8.2)
head('4.2 Synthetic feature ablation',2)
ablrows=[]
for a in syn['ablation_summary']:
 ablrows.append([a['ablation'].replace('_',' '),f"{a['mae_mean']:.4f} ± {a['mae_sd']:.4f}",f"{a['rmse_mean']:.4f} ± {a['rmse_sd']:.4f}"])
table('Table 4. Synthetic LightGBM ablation performance by fold (mean ± SD)',['Feature set','MAE (project units)'],[[r[0],r[1]] for r in ablrows])
testrows=[]
for z in syn['paired_daily_effects']:
 intervals=z['moving_block_bootstrap_95ci_by_days']
 block_text='; '.join(f"{length}d [{intervals[str(length)][0]:.4f}, {intervals[str(length)][1]:.4f}]" for length in (7,14,30))
 testrows.append([z['comparison'].replace('without_','Without ').replace('_',' '),f"{z['mean_daily_mae_delta']:.4f}; {block_text}"])
table('Table 5. Date-paired synthetic ablation effects and block-length sensitivity',['Ablation','Mean ΔMAE and 95% intervals by block length'],testrows)
for x in [
"The full-feature LightGBM model had a mean fold MAE of 0.4674 project-defined tonnes-equivalent (fold standard deviation 0.1460). Removing weather/calendar variables increased date-averaged MAE by 0.0569 units (14-day block interval 0.0389 to 0.0769); removing target history increased it by 0.0508 (0.0180 to 0.0793). For both comparisons, intervals remained above zero across 7-, 14- and 30-day blocks. These are exploratory estimates conditional on this simulator and evaluation period.",
"Removing production increased daily MAE by 0.0063 units; its intervals included zero for 14- and 30-day blocks, while the 7-day interval was narrowly positive. Removing price and operating-status features increased MAE by 0.0047, with all three block-length intervals including zero. The estimates do not establish that these inputs are generally unhelpful. Their programmed variation may have been weak relative to target persistence.",
"Ablation changes the model’s available information and may also alter how the learner partitions the remaining variables. Consequently, the differences are conditional on this model, tuning, data-generating process and forecast horizon. They do not show that manipulating weather or recent energy consumption changes actual industrial demand, nor do they establish confirmatory evidence beyond this synthetic benchmark."
]:para(x)
fig(OUT/'figures'/'ablation_mae.png','Figure 2. Fold-level MAE across full and ablated synthetic models.',width=8.2)
head('4.3 Synthetic TreeSHAP attribution',2)
for x in [
"In the first synthetic test fold, tce_lag_1 had the largest mean absolute SHAP value (3.873 project units), followed by the seven-day rolling mean (2.069), the 14-day lag (1.749), the 28-day rolling mean (1.366) and the seven-day lag (1.254). The remaining individual inputs had substantially smaller mean absolute contributions in this 60-case sample. The prominence of recent energy history agrees with the ablation result at the feature-group level.",
"This agreement should be interpreted narrowly. The lag and moving-average features are correlated by construction, and the simulator contains persistence. TreeSHAP can distribute contribution among these related inputs. Meanwhile, weather/calendar ablation has a measurable effect even though the individual importance is distributed across temperature, weekday and annual encodings. Group ablation and single-feature attribution answer different questions: the former removes a set jointly, while the latter allocates fitted prediction differences among the model inputs."
]:para(x)
fig(OUT/'figures'/'shap_importance.png','Figure 3. Mean absolute TreeSHAP values in the first synthetic test fold.',width=8.2)

head('5. Public steel-facility application',1)
for x in [
"The UCI data were aggregated from 35,040 quarter-hour measurements into 365 complete daily electricity totals. The real-data evaluation used seven rolling folds, with 30 test days per fold and a 120-day initial training period after a 28-day warm-up. LightGBM achieved pooled MAE of 769.9 kWh/day and RMSE of 1,039.2 kWh/day over 210 test days. The seven-day seasonal-naive baseline achieved MAE of 869.6 kWh/day and RMSE of 1,167.5 kWh/day. LightGBM therefore reduced pooled MAE by 99.8 kWh/day, or 11.5%, relative to the baseline. The fold series shows that model performance was not uniform: the largest LightGBM MAE occurred in the seventh fold, at 1,220.8 kWh/day.",
"The public-series ablations provide a qualified picture. Removing recent history increased date-paired MAE by 152.0 kWh/day; the 14-day moving-block interval was 78.7 to 216.6, and the interval remained positive for 7- and 30-day blocks. Removing monthly history increased MAE by 49.5 kWh/day, but its 30-day block interval included zero (−4.2 to 84.2). Removing calendar variables reduced MAE by 34.2 kWh/day; all block-length intervals included zero. With one facility-year, these estimates are exploratory and do not establish that calendar variables are harmful or that long-term lags lack value.",
"The first-fold TreeSHAP ranking placed day-of-week sine encoding highest (mean absolute attribution 795.5 kWh/day), followed by the one-day lag (296.2), seven-day mean (200.9), annual day cosine (160.3) and weekend indicator (134.4). The attribution ranking and date-paired ablation answer different questions: one describes a fitted first-fold model and the other compares loss over the full test period."
]:para(x)
table('Table 6. Public steel-facility forecast results',['Model / ablation','MAE or ΔMAE; block inference'],[
['LightGBM','769.9 kWh/day; 210 test days'],['7-day seasonal naive','869.6 kWh/day; 210 test days'],['Without recent history','+152.0 kWh/day; 14d CI [78.7, 216.6]'],['Without monthly history','+49.5 kWh/day; 30d CI includes zero'],['Without calendar','−34.2 kWh/day; all CIs include zero']])
fig(REAL/'../figures'/'uci_steel'/'uci_fold_mae.png','Figure 4. Fold-level MAE on the public steel-facility series.',width=8.2)
fig(REAL/'../figures'/'uci_steel'/'uci_shap_importance.png','Figure 5. Mean absolute TreeSHAP values in the first public-data test fold.',width=8.2)

head('6. Discussion',1)
head('6.1 Interpretation across the two data settings',2)
for x in [
"Both data settings indicate that recent target history is a useful predictive input for next-day aggregate energy use under a rolling protocol. This result is plausible for daily energy series with persistence, but the two estimates have different meanings. In the synthetic panel, the relation is partly imposed by the simulator and shared seasonal structure. In the UCI data, it is observed in one facility across one calendar year. Neither result shows that a historical lag is a causal driver or that the same predictor will perform at another plant.",
"Calendar and weather features behaved differently across settings. The synthetic generator explicitly includes seasonality and temperature-related variation, and weather/calendar removal increased date-paired error across tested block lengths. In the measured steel series, calendar variables had high local SHAP attribution but their date-paired effect intervals included zero. The difference may reflect the simulated design, the real facility’s operating schedule, the short evaluation horizon or sampling variability. The data do not allow these explanations to be separated.",
"LightGBM outperformed the baseline and sequence models in pooled MAE on the synthetic panel; it also outperformed the seasonal-naive baseline in the UCI application. This is evidence about the tested configurations, features and folds. It is not evidence that gradient boosting is universally preferable to neural networks. The LSTM and Transformer were not exhaustively tuned, and the small number of training years gives limited support for estimating the value of complex sequence architectures. A future benchmark should report tuning budgets, computational cost and performance across independent plants, not only a larger model catalogue."
]:para(x)
head('6.2 Methodological and operational implications',2)
for x in [
"For a real forecasting deployment, the first requirement is a timestamped data dictionary that states when each feature becomes available. Production plans, actual output, weather forecasts and observed temperatures may have different availability. An offline model that uses a field recorded after the forecast origin can appear accurate while being unusable in operation. The explicit one-day lag rule used here is a conservative device; it should be replaced by field-specific availability checks when operational data are available.",
"The second requirement is a forecast benchmark that is hard to beat. A seven-day lag can be highly competitive for stable weekly patterns and makes a useful check against complex models. Rolling windows provide a view of temporal variation that a random split conceals. If the intended use is a weekly energy plan, the forecast horizon and evaluation unit should match that decision. This paper evaluates only one-day-ahead aggregate energy and does not assess monthly procurement, intraday scheduling or energy-intensity forecasting.",
"The third requirement is uncertainty reporting at an appropriate replication level. With few rolling folds, conclusions should be presented as exploratory estimates accompanied by intervals and fold-level plots. Daily test points within a 30-day block are not independent experimental replications. For multi-facility claims, the independent unit should include facilities or operating periods, ideally with a held-out plant and a longer historical span. Additional data would also allow evaluation of recalibration, structural breaks and model drift.",
"A useful extension would report performance under more than one forecast horizon. A one-day horizon emphasizes short-term persistence and may reward recent lags. A week-ahead or month-ahead schedule would reduce the usefulness of some lag values unless future predictions are generated recursively, in which case errors accumulate. Horizon-specific benchmarks would need a decision-aligned loss function, since underprediction and overprediction may carry different operational costs. The current MAE and RMSE treat both directions symmetrically and measure energy-unit error; they do not encode tariff periods, demand charges, production shortfalls or emissions intensity.",
"The evaluation could also distinguish between refitting and updating. In this study, model training expands at each fold, while features are recalculated from the available history. A live system might retrain on a fixed schedule, update daily, or preserve a model through a production campaign. Each choice changes how quickly the forecast can adapt and how much recent data influences it. Prospective shadow deployment would therefore be needed before operational adoption: predictions should be logged at their true issue time, compared with later observations, and reviewed across operating states before any control action is connected.",
"Evaluation across operating subgroups would be another important extension. Overall workshop-day MAE can hide poor forecasts for a particular workshop, energy carrier, weekend, shutdown period or high-load state. The synthetic panel contains some of these labels, but subgroup comparisons would remain conditional on generated entities and schedules. The UCI data have a single facility and daily target, limiting meaningful process-specific strata. A future study with authorized facility data should prespecify relevant strata, report sample counts alongside errors, and avoid interpreting small subgroup differences when the number of periods is low.",
"Finally, model explanations should support inspection rather than replace engineering evidence. A SHAP ranking can help identify influential model inputs and cases for review. It cannot establish whether a control action will reduce consumption, explain an abnormal event, or assign responsibility for a deviation. Such claims require validated process measurements, documented interventions or confirmed incident labels. No anomaly detection or fault diagnosis is evaluated here."
]:para(x)
head('6.3 Limitations',2)
for x in [
"The synthetic study has limited external validity by design. The generator includes explicit trend, seasonality, production and noise assumptions. Even when it injects dirty fields and missingness, it does not reproduce the full uncertainty of plant instrumentation, product mix, process constraints, maintenance or operator decisions. The synthetic panel is useful for checking code paths and inference procedures but not for estimating field error or energy-saving benefit.",
"The public-data application covers one steel facility and one year. Daily aggregation reduces 35,040 measurements to 365 cases and removes within-day variation. No production quantity or process metadata is available, and one facility cannot represent between-plant variation. The modest advantage over the baseline may also depend on the selected evaluation blocks. A longer record would be needed to assess interannual seasonality, structural changes and performance under different production regimes.",
"The data ownership and licence also define the limits of reuse. The UCI repository makes this benchmark available under CC BY 4.0, which permits reuse with appropriate attribution. It does not provide permission to access a private factory historian or to publish sensitive operational data from another facility. A field replication should document data authorization, confidentiality arrangements, aggregation rules and the exact version of the extraction used. For reproducibility, the present analysis records dataset identity and a file hash in the project provenance metadata; readers can use the repository citation to retrieve the public source and compare versions. The project coefficients should be reconciled against the applicable requirements of GB/T 2589-2020 before any real-facility total-energy claim is made [30].",
"The inferential analysis uses 330 paired synthetic dates and 210 paired public-data dates, but they come from one generated panel and one facility-year. Block-bootstrap intervals vary with block length and cannot establish between-facility or between-year generality. All ablation estimates are exploratory; they are not confirmatory hypothesis tests.",
"The explanation analysis is similarly limited to one fitted fold model in each setting and 60 or 30 cases. Feature correlation affects attribution, while importance can change as the model is retrained. We have not calculated SHAP stability across folds, subgroups or alternative backgrounds. The study also does not quantify forecast calibration, prediction intervals, economic costs, emissions effects or the consequences of acting on a forecast."
,
"The results should also be read in relation to aggregation. Daily totals suppress within-day peaks and ramps that can determine electrical demand charges or constrain equipment operation. The method could be adapted to hourly or quarter-hour targets, but that would require a redesigned validation scheme, attention to intra-day leakage, and baselines matched to shorter seasonal cycles. We do not infer sub-daily behaviour from the daily UCI totals."
]:para(x)

head('7. Conclusion',1)
for x in [
"This study evaluated a next-day industrial energy forecasting workflow with expanding-window validation, date-paired ablation and TreeSHAP attribution. On the fixed-seed synthetic panel, LightGBM achieved lower pooled MAE than the tested seasonal-naive, LSTM and Transformer configurations. Removing weather/calendar or target-history features increased date-averaged MAE, with block-bootstrap intervals positive across the tested block lengths. Recent target lags dominated the first-fold SHAP ranking. These results remain conditional on the constructed panel.",
"Applied separately to a public steel-facility electricity series, LightGBM improved on the seven-day baseline in pooled MAE, and removing recent history increased date-paired error across block lengths. The calendar-group effect intervals included zero. Neither experiment establishes general performance across industrial facilities; the synthetic target is a project-defined mixed resource index, and the public series covers one site and one year.",
"The results support a reproducible method for temporal evaluation and model inspection. They do not establish general performance across industrial facilities, causal effects of features, energy savings or fault-warning capability. The next empirical step is validation on longer, authorized multi-facility data with production quantities, explicit feature timestamps and independently confirmed operational events."
]:para(x)

head('Acknowledgements',1);para('No external funding or assistance is declared in this draft. The authors should confirm or replace this statement before submission.')
head('Nomenclature',1)
for x in ['MAE  mean absolute error','RMSE  root mean squared error','SHAP  SHapley Additive exPlanations','project-equivalent unit  synthetic resource-index unit derived from project conversion factors','UCI  University of California, Irvine Machine Learning Repository','y(t−k)  energy observation k days before target date t']:
 para(x)
head('References',1)
raw=(ROOT/'docs'/'论文参考文献.md').read_text(encoding='utf-8').splitlines()
ref_lines=[line for line in raw if re.match(r'^\d+\. ',line)]
# Cite in order of first appearance in the manuscript: rolling-origin sources,
# forecasting text, predictive-accuracy test, SHAP, adjacent reviews, then data.
ref_by_number={int(re.match(r'^(\d+)\.',line).group(1)):line for line in ref_lines}
ref_order=list(range(1,22))+[24,25,23,22,26,27,28,29,30]
ref_lines=[ref_by_number[i] for i in ref_order]
for i,line in enumerate(ref_lines,1):
 para(f'[{i}] '+re.sub(r'^\d+\.\s*','',line).replace('*',''))

# Page number fields in footer
for s in doc.sections:
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
doc.save(DEST)
words=sum(len(re.findall(r"\b[\w−'-]+\b",p.text)) for p in doc.paragraphs)
print(f'{DEST}\nword_like_count={words}\nreferences={len(ref_lines)}')
